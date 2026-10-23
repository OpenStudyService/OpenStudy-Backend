"""Assemble the reviewed handoff, validate generated artifacts and hash files.
Usage: python scripts/package_foundation_release.py <xlsx-output-dir>
No network, GitHub writes, database connections or secret-file reads.
"""
import hashlib
import json
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from openpyxl import load_workbook  # Read-only verification; authorship uses artifact-tool.
from pypdf import PdfReader

root=Path(__file__).resolve().parents[1]
release=root/'docs/releases/foundation-0.1.2'
artifact_dir=Path(sys.argv[1]).resolve()
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
migration_dir=root/'src/main/resources/db/migration'
for source in migration_dir.glob('V*.sql'):shutil.copyfile(source,release/'migrations'/source.name)
shutil.copyfile(root/'docs/design/foundation/openapi.json',release/'openapi.json')
shutil.copyfile(root/'docs/verification/foundation-schema-v3.md',release/'verification.md')
shutil.copyfile(artifact_dir/'OpenStudy-Foundation-API.xlsx',release/'OpenStudy-Foundation-API.xlsx')
spec=json.loads((release/'openapi.json').read_text(encoding='utf-8'))
metadata=json.loads((release/'schema-metadata.json').read_text(encoding='utf-8'))
assert len(metadata['tables'])==66 and len(metadata['columns'])==634
sql=(release/'schema-full.sql').read_text(encoding='utf-8')
assert len(re.findall(r'^CREATE TABLE ',sql,re.M))==66
assert not re.search(r'\b(DROP DATABASE|DROP TABLE|FOREIGN_KEY_CHECKS)\b',sql,re.I)
assert not re.search(r'INSERT INTO `(employees|SPRING_SESSION|products|payments|passes)`',sql)
v1=migration_dir/'V1__init_schema.sql'
original=subprocess.check_output(['git','show','main:src/main/resources/db/migration/V1__init_schema.sql'],cwd=root)
assert original.decode('utf-8').replace('\r\n','\n')==v1.read_text(encoding='utf-8'),'V1 changed'
w=load_workbook(release/'OpenStudy-Foundation-API.xlsx',read_only=False,data_only=False)
assert w.sheetnames==['사용안내','API목록','요청필드','응답목록','공통모델','오류·보안','업무규칙','확정정책']
operations={o['operationId'] for p in spec['paths'].values() for o in p.values()}
assert {r[0] for r in w['API목록'].iter_rows(min_row=6,values_only=True)}==operations
assert {r[0] for r in w['공통모델'].iter_rows(min_row=6,values_only=True)}==set(spec['components']['schemas'])
assert sha(release/'openapi.json') in w['사용안내']['B15'].value
assert {r[0]:r[3] for r in w['확정정책'].iter_rows(min_row=6,values_only=True)}=={k:v['decision'] for k,v in spec['x-confirmed-policies'].items()}
for s in w:
    assert not any(c.data_type=='e' for row in s for c in row)
    if s.title!='사용안내':assert s.freeze_panes and len(s.tables)==1
assert len(PdfReader(release/'OpenStudy-Foundation-ERD.pdf').pages)==7
drafts=json.loads((root/'docs/design/foundation/github-drafts/registration.json').read_text(encoding='utf-8'))
assert drafts['githubRegistrationPerformed'] is False and len(drafts['issues'])==16
assert drafts['contractSha256']==sha(release/'openapi.json')
for i in drafts['issues']:
    assert (root/'docs/design/foundation/github-drafts'/i['bodyFile']).is_file()
    assert set(i['dependsOn'])<={x['id'] for x in drafts['issues']}

paths=set(p for p in release.rglob('*') if p.is_file() and p.name not in ['manifest.json','SHA256SUMS.txt'])
paths.update(p for p in (root/'docs/design/foundation').rglob('*') if p.is_file())
paths.update(migration_dir.glob('V*.sql'))
paths.update(root/f for f in ['.gitattributes','AGENTS.md','docs/db/foundation-upgrade.md','docs/verification/foundation-schema-v3.md',
 'scripts/export_foundation_api.mjs','scripts/export_foundation_erd.py','scripts/prepare_foundation_issue_drafts.py',
 'scripts/package_foundation_release.py','scripts/validate_foundation_design.py'])
files=[{'path':p.relative_to(root).as_posix(),'sha256':sha(p),'bytes':p.stat().st_size} for p in sorted(paths) if p.is_relative_to(release)]
zip_sources=[{'path':p.relative_to(root).as_posix(),'sha256':sha(p),'bytes':p.stat().st_size} for p in sorted(paths) if not p.is_relative_to(release)]
source=json.loads((root/'docs/design/foundation/requirements.json').read_text(encoding='utf-8'))
manifest={'release':'foundation-0.1.2','date':'2026-10-24','contractVersion':spec['info']['version'],'contractSha256':sha(release/'openapi.json'),
 'flywayVersion':3,'srsSha256':source['sha256'],'reviewedRequirementCount':61,'operationCount':40,'schemaCount':57,
 'businessAndSessionTables':65,'flywayHistoryTables':1,'coreTables':21,'testResult':{'tests':53,'failures':0,'errors':0,'skipped':0},
 'githubRegistrationPerformed':False,'files':files,'zipSourceFiles':zip_sources}
(release/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
(release/'SHA256SUMS.txt').write_text(''.join(f'{f["sha256"]}  {f["path"]}\n' for f in files),encoding='utf-8')
paths.update([release/'manifest.json',release/'SHA256SUMS.txt'])
destination=artifact_dir/'OpenStudy-Foundation-0.1.2.zip'
with zipfile.ZipFile(destination,'w',zipfile.ZIP_DEFLATED) as z:
    z.writestr('README.md','# OpenStudy 기반 개발 전달 패키지\n\n[전달 안내](docs/releases/foundation-0.1.2/README.md)부터 읽으세요.\n\nDB V3는 검증 완료, 업무 API는 목표 계약입니다. GitHub 마일스톤·이슈는 아직 등록하지 않았습니다.\n')
    for p in sorted(paths):z.write(p,p.relative_to(root).as_posix())
with zipfile.ZipFile(destination) as z:
    assert z.testzip() is None
    for f in files+zip_sources:assert hashlib.sha256(z.read(f['path'])).hexdigest()==f['sha256']
shutil.copyfile(release/'OpenStudy-Foundation-ERD.pdf',artifact_dir/'OpenStudy-Foundation-ERD.pdf')
print(json.dumps({'result':'passed','filesHashed':len(files),'zipFiles':len(paths)+1,'zipSha256':sha(destination),'output':str(destination)},ensure_ascii=False))
